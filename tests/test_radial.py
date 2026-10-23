"""The tools menu (app/radial): every tool's own work, without the screen.

Zip and extract, media out of a deck and a Word file, video to MP4,
PDF merge/split/lock, the scrolling-capture stitch, the GIF, the calendar,
the world clock against the real time-zone database, the QR encoder, the
icons, file transfer end to end over the loopback, and that every tool
saves into its own folder in Downloads/Squint.

Run before every release (see .github/workflows/release.yml)."""
import datetime as dt
import json
import os
import shutil
import socket
import sys
import tempfile
import time
import unittest
import urllib.error
import urllib.request
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

TMP = tempfile.mkdtemp(prefix="squint-radial-")
# Everything the tools write (Downloads/Squint, settings) lands in here.
os.environ["HOME"] = TMP
os.environ["USERPROFILE"] = TMP
os.environ["XDG_CONFIG_HOME"] = TMP

from PIL import Image, ImageDraw                               # noqa: E402

from app.radial import (calendar_view as cal, capture, common, docs,  # noqa
                        icons, qr, record, send, worldclock)


def _has(module):
    try:
        __import__(module)
        return True
    except Exception:
        return False


def _picture(name="p.png", size=(120, 80), colour="red"):
    path = os.path.join(TMP, name)
    Image.new("RGB", size, colour).save(path)
    return path


class Folders(unittest.TestCase):
    def test_each_tool_has_its_folder(self):
        for tool in ("shot", "rec", "pdf", "zip", "send", "ext"):
            path = common.squint_dir(tool)
            self.assertTrue(os.path.isdir(path), path)
            self.assertEqual(os.path.dirname(path),
                             os.path.join(TMP, "Downloads", "Squint"))
        z = common.into("zip", "a.zip")
        self.assertIn(os.path.join("Downloads", "Squint", "Zip"), z)
        self.assertEqual(common.where("/somewhere/else.txt"), "")


class Zip(unittest.TestCase):
    def test_zip_and_extract(self):
        src = os.path.join(TMP, "src")
        os.makedirs(os.path.join(src, "sub"), exist_ok=True)
        with open(os.path.join(src, "a.txt"), "w") as f:
            f.write("hello")
        with open(os.path.join(src, "sub", "b.txt"), "w") as f:
            f.write("world")
        z = docs.zip_plain([src], os.path.join(TMP, "x.zip"))
        self.assertFalse(docs._needs_password(z))
        out = docs.extract_plain(z, os.path.join(TMP, "unz"))
        with open(os.path.join(out, "src", "sub", "b.txt")) as f:
            self.assertEqual(f.read(), "world")


class Media(unittest.TestCase):
    def test_from_powerpoint(self):
        deck = os.path.join(TMP, "deck.pptx")
        with zipfile.ZipFile(deck, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml", "<x/>")
            zf.write(_picture(), "ppt/media/image1.png")
            zf.writestr("ppt/media/media2.mp4", b"\0" * 100)
        folder, n = docs.extract_media(deck, TMP)
        self.assertEqual(n, 2)
        self.assertEqual(sorted(os.listdir(folder)),
                         ["image1.png", "media2.mp4"])

    @unittest.skipUnless(_has("docx"), "python-docx not installed")
    def test_from_word(self):
        import docx
        d = docx.Document()
        d.add_paragraph("hi")
        d.add_picture(_picture())
        path = os.path.join(TMP, "doc.docx")
        d.save(path)
        _folder, n = docs.extract_media(path, TMP)
        self.assertEqual(n, 1)


class Video(unittest.TestCase):
    def test_formats(self):
        for ext in (".mov", ".mkv", ".avi", ".wmv", ".webm", ".m4v", ".mts"):
            self.assertIn(ext, docs.VIDEO_EXT)

    @unittest.skipUnless(shutil.which("ffmpeg"), "FFmpeg not installed")
    def test_to_mp4(self):
        src = os.path.join(TMP, "clip.mkv")
        common.run([shutil.which("ffmpeg"), "-loglevel", "error", "-y", "-f",
                    "lavfi", "-i", "testsrc=size=160x120:rate=10", "-t",
                    "1", src])
        out = docs.video_to_mp4(shutil.which("ffmpeg"), src)
        self.assertTrue(out.endswith(".mp4"))
        self.assertIn(os.path.join("Squint", "Recordings"), out)


class Tools(unittest.TestCase):
    def test_eight_tools_no_convert(self):
        keys = [t["key"] for t in __import__("app.radial",
                                             fromlist=["_tools"])._tools()]
        self.assertEqual(keys, ["shot", "rec", "trans", "pdf", "zip", "send",
                                "cal", "ext"])


@unittest.skipUnless(_has("pymupdf") or _has("fitz"), "PyMuPDF not installed")
class Pdf(unittest.TestCase):
    def _pdf(self, name, pages):
        fitz = docs._fitz()
        doc = fitz.open()
        for i in range(pages):
            doc.new_page().insert_text((72, 72), "%s page %d" % (name, i + 1))
        path = os.path.join(TMP, name + ".pdf")
        doc.save(path)
        doc.close()
        return path

    def test_merge_split_lock(self):
        fitz = docs._fitz()
        a, b = self._pdf("a", 2), self._pdf("b", 3)
        merged = docs.merge_pdfs([a, b], os.path.join(TMP, "m.pdf"))
        with fitz.open(merged) as doc:
            self.assertEqual(doc.page_count, 5)
        folder = docs.split_pdf(merged)
        self.assertEqual(len(os.listdir(folder)), 5)
        locked = docs.lock_pdf(a, os.path.join(TMP, "l.pdf"), "s3cret")
        with fitz.open(locked) as doc:
            self.assertTrue(doc.needs_pass)
            self.assertTrue(doc.authenticate("s3cret"))
        with self.assertRaises(RuntimeError):
            docs.merge_pdfs([locked], os.path.join(TMP, "x.pdf"))

    def test_files_to_pdf(self):
        out = docs.files_to_pdf([_picture("f1.png"), _picture("f2.png")],
                                os.path.join(TMP, "pics.pdf"))
        with docs._fitz().open(out) as doc:
            self.assertEqual(doc.page_count, 2)


class Capture(unittest.TestCase):
    def test_scroll_stitch(self):
        page = Image.new("RGB", (400, 2400), "white")
        draw = ImageDraw.Draw(page)
        for y in range(0, 2400, 30):
            n = y // 30
            draw.text((10 + n % 50, y + 5), "Line %d %s" % (n, "x" * (n % 17)),
                      fill="black")
            draw.rectangle((300, y + 4, 300 + (y * 7) % 90, y + 20),
                           fill=(y % 255, 80, 160))
        frames, pos, h = [], 0, 600
        while pos + h <= 2400:
            frames.append(page.crop((0, pos, 400, pos + h)))
            pos += 137
        frames.append(frames[-1])         # the page stopped moving
        tall = capture.stitch(frames)
        self.assertEqual(tall.height, (len(frames) - 2) * 137 + h)
        self.assertEqual(tall.tobytes(),
                         page.crop((0, 0, 400, tall.height)).tobytes())

    def test_gif(self):
        frames = []
        for i in range(20):
            frame = Image.new("RGB", (200, 120), "white")
            ImageDraw.Draw(frame).rectangle((i * 8, 30, i * 8 + 30, 90),
                                            fill="blue")
            frames.append(frame)
        path = record.save_gif(frames, TMP)
        with Image.open(path) as gif:
            self.assertEqual(gif.n_frames, 20)


class Calendar(unittest.TestCase):
    def test_agenda(self):
        day = dt.date(2026, 10, 22)
        cal.add(day, "Design review", "10:00")
        cal.add(day, "Lunch")
        cal.add(day, "Call", "09:00")
        items = cal.load()[day.isoformat()]
        self.assertEqual([i["time"] for i in items], ["09:00", "10:00", ""])
        cal.remove(day, 0)
        self.assertEqual(len(cal.load()[day.isoformat()]), 2)
        self.assertEqual(cal.parse_time("930"), "09:30")
        self.assertEqual(cal.parse_time("14.15"), "14:15")
        self.assertIsNone(cal.parse_time("25"))


class WorldClock(unittest.TestCase):
    def test_built_in_rules_match_the_time_zone_database(self):
        try:
            from zoneinfo import ZoneInfo
            ZoneInfo("Europe/Berlin")
        except Exception:
            self.skipTest("no time-zone database here")
        start = dt.datetime(2025, 1, 1, tzinfo=dt.timezone.utc)
        for name, (zone, std, rule) in worldclock.CITIES.items():
            for hours in range(0, 3 * 365 * 24, 7):
                utc = start + dt.timedelta(hours=hours)
                real = utc.astimezone(ZoneInfo(zone))
                naive = utc.replace(tzinfo=None)
                summer = worldclock._in_dst(rule, naive, std)
                ours = naive + dt.timedelta(minutes=std + (60 if summer else 0))
                self.assertEqual(ours, real.replace(tzinfo=None),
                                 "%s at %s" % (name, utc))

    def test_defaults_exist(self):
        for name in worldclock.DEFAULT:
            self.assertIn(name, worldclock.CITIES)


class QrAndIcons(unittest.TestCase):
    def test_qr_encodes(self):
        for text in ("https://www.paypal.com/qrcodes/p2pqrc/VCMMGZLN8W22C",
                     "http://192.168.1.20:53317/?k=" + "x" * 32):
            grid = qr.encode(text)
            size = len(grid)
            self.assertEqual((size - 17) % 4, 0)
            self.assertTrue(all(len(row) == size for row in grid))
        if _has("pyzbar"):
            from pyzbar.pyzbar import decode
            img = qr.render_round(text, 600, "#FFFFFF")
            self.assertEqual(decode(img)[0].data.decode(), text)

    def test_icons(self):
        for key in ("shot", "rec", "trans", "pdf", "zip", "send", "cal",
                    "ext"):
            img = icons.image(key, 22, "#FFFFFF")
            self.assertEqual(img.size, (22, 22))
            self.assertGreater(img.getchannel("A").getbbox()[2], 10, key)


class Transfers(unittest.TestCase):
    def test_end_to_end(self):
        shared = os.path.join(TMP, "share.bin")
        with open(shared, "wb") as f:
            f.write(os.urandom(1024 * 1024 + 7))
        pic = _picture("t.png")
        received = common.received_dir()
        for old in os.listdir(received):
            os.remove(os.path.join(received, old))
        srv = send.Server([shared], receive=True)
        try:
            srv.ask = lambda offer: True
            base = "http://127.0.0.1:%d" % srv.port
            op = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            page = op.open(base + "/?k=" + srv.key).read().decode()
            self.assertIn("share.bin", page)
            with open(shared, "rb") as f:
                self.assertEqual(op.open(base + "/f/0?k=" + srv.key).read(),
                                 f.read())
            with self.assertRaises(urllib.error.HTTPError) as err:
                op.open(base + "/?k=wrong")
            self.assertEqual(err.exception.code, 403)
            # A name trying to climb out of the folder stays inside it.
            req = urllib.request.Request(
                base + "/put?k=%s&name=..%%2F..%%2Fevil.txt" % srv.key,
                data=b"phone data", method="POST")
            self.assertEqual(op.open(req).status, 200)
            self.assertEqual(os.listdir(received), ["evil.txt"])
            n = send.push({"ip": "127.0.0.1", "port": srv.port,
                           "name": "me"}, [shared, pic])
            self.assertEqual(n, 2)
            self.assertEqual(sorted(os.listdir(received)),
                             ["evil.txt", "share.bin", "t.png"])
            srv.ask = lambda offer: False
            with self.assertRaises(RuntimeError):
                send.push({"ip": "127.0.0.1", "port": srv.port,
                           "name": "me"}, [pic])
            time.sleep(0.3)
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.settimeout(3)
            s.sendto(send.HELLO, ("127.0.0.1", send.DISCOVERY_PORT))
            self.assertEqual(json.loads(s.recvfrom(1024)[0])["port"],
                             srv.port)
            s.close()
        finally:
            srv.stop()


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
